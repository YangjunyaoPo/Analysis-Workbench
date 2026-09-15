"""Inspect downloaded tables and create controlled curve fixtures for discovery.

Requires numpy, openpyxl, reportlab and pypdfium2 in the analysis environment.
The generated pictures are rendered from published numeric data, not instrument exports.
"""

import csv
import hashlib
import json
import math
from pathlib import Path
import subprocess

import numpy as np
import openpyxl
import pypdfium2 as pdfium
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics import renderPDF
from reportlab.lib import colors


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "initial-samples"


def summarize_matrix(x, y):
    return {
        "coordinate_count": len(x),
        "coordinate_range": [float(x.min()), float(x.max())],
        "coordinate_step_range": [float(np.diff(x).min()), float(np.diff(x).max())],
        "coordinates_strictly_increasing": bool(np.all(np.diff(x) > 0)),
        "signal_shape": list(y.shape),
        "nonfinite_signal_values": int((~np.isfinite(y)).sum()),
        "rows_with_nonfinite_signal": int((~np.isfinite(y)).any(axis=1).sum()),
        "coordinates_with_nonfinite_signal": int((~np.isfinite(y)).any(axis=0).sum()),
        "negative_signal_values": int((y < 0).sum()),
        "zero_signal_values": int((y == 0).sum()),
        "signal_range": [float(np.nanmin(y)), float(np.nanmax(y))],
    }


def render_curve(points, filename, sample_id):
    drawing = Drawing(720, 450)
    plot = LinePlot()
    plot.x, plot.y, plot.width, plot.height = 75, 75, 595, 305
    plot.data = [points]
    plot.lines[0].strokeColor = colors.HexColor("#1261A0")
    plot.lines[0].strokeWidth = 1.2
    plot.xValueAxis.valueMin, plot.xValueAxis.valueMax = 1600, 1800
    plot.xValueAxis.valueSteps = list(range(1600, 1801, 50))
    plot.yValueAxis.valueMin, plot.yValueAxis.valueMax = 0, 0.1
    plot.yValueAxis.valueSteps = [0, 0.02, 0.04, 0.06, 0.08, 0.1]
    drawing.add(plot)
    drawing.add(String(75, 416, f"FTIR reference curve: {sample_id}", fontSize=15))
    drawing.add(String(75, 396, "Generated from published numeric data; not an instrument image", fontSize=10))
    drawing.add(String(270, 34, "Wavenumber (cm^-1)", fontSize=11))
    drawing.add(String(10, 265, "Signal", fontSize=10))
    drawing.add(String(75, 12, "Source: Jimenez-Carvelo, Zenodo 14651816, CC BY 4.0", fontSize=9))
    document = pdfium.PdfDocument(renderPDF.drawToString(drawing))
    page = document[0]
    bitmap = page.render(scale=2)
    bitmap.to_pil().save(filename)
    bitmap.close()
    page.close()
    document.close()


def main():
    reports = []
    uv_path = DATA / "19324549" / "Chemcial_Data_for_PLS.csv"
    with uv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    x = np.array(rows[0][9:], dtype=float)
    y = np.array([row[9:] for row in rows[1:]], dtype=float)
    reports.append({"file": uv_path.name, **summarize_matrix(x, y),
                    "metadata_columns": rows[0][:9],
                    "coordinate_units": "Not explicit in CSV header; requires source confirmation",
                    "experiment_dates": "No explicit date column"})

    folder = DATA / "14651816"
    members = json.loads((folder / "archive-members.json").read_text())
    extracted = folder / "extracted"
    extracted.mkdir(exist_ok=True)
    for member in members:
        if not member.endswith(".xlsx"):
            continue
        path = extracted / Path(member).name
        if not path.exists():
            body = subprocess.run(
                ["tar", "-xOf", str(folder / "Data_set_FTIR_Raman.rar"), member],
                capture_output=True, check=True, timeout=30,
            ).stdout
            path.write_bytes(body)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        x = np.array(rows[0][1:], dtype=float)
        y = np.array([row[1:] for row in rows[1:]], dtype=float)
        ids = [row[0] for row in rows[1:]]
        result = {
            "file": path.name, "sheet": sheet.title,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "id_header": rows[0][0], "sample_count": len(ids),
            "duplicate_ids": len(ids) - len(set(ids)),
            "first_sample_ids": ids[:5],
            "coordinate_units": "cm^-1 according to linked paper instrumentation section",
            "experiment_dates": "Not present in these sheets",
            **summarize_matrix(x, y),
        }
        if path.name == "FTIR_spectra_Classification Study.xlsx":
            selected = (x >= 1600) & (x <= 1800)
            band_x, band_y = x[selected], y[0, selected]
            points = [(float(a), float(b)) for a, b in zip(band_x, band_y)]
            area_numpy = float(np.trapezoid(band_y, band_x))
            area_reference = math.fsum(
                (x1 - x0) * (y0 + y1) / 2
                for (x0, y0), (x1, y1) in zip(points, points[1:])
            )
            peak_index = int(np.argmax(band_y))
            result["reference_example"] = {
                "sample_id": ids[0], "sheet_row": 2,
                "requested_coordinate_interval": [1600, 1800],
                "actual_sampled_interval": [points[0][0], points[-1][0]],
                "points": len(points),
                "largest_sampled_value_coordinate": float(band_x[peak_index]),
                "largest_sampled_value": float(band_y[peak_index]),
                "trapezoidal_area_numpy": area_numpy,
                "trapezoidal_area_independent_scalar": area_reference,
                "absolute_difference": abs(area_numpy - area_reference),
                "interpretation": "Raw signal area; no baseline correction, peak fitting, or concentration inference",
            }
            fixture = DATA / "controlled-fixtures"
            fixture.mkdir(exist_ok=True)
            (fixture / "ftir-reference.json").write_text(json.dumps({
                "origin": "Published numeric measurements; locally rendered picture",
                "source_record": "https://zenodo.org/records/14651816",
                "creator": "Ana M. Jimenez-Carvelo",
                "license": "https://creativecommons.org/licenses/by/4.0/",
                "changes": "First sample selected, coordinate interval cropped, picture rendered",
                "sample_id": ids[0], "points": points,
                "source_sha256": result["sha256"],
            }, indent=2), encoding="utf-8")
            render_curve(points, fixture / "ftir-reference.png", ids[0])
        reports.append(result)
        print(json.dumps(result, indent=2), flush=True)
    (DATA / "table-analysis.json").write_text(json.dumps(reports, indent=2), encoding="utf-8")
    print(json.dumps(reports[0], indent=2))


if __name__ == "__main__":
    main()
